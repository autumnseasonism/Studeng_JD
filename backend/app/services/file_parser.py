import os
from pathlib import Path

def parse_file(file_path: str) -> dict:
    """解析文件并返回文本内容
    
    Args:
        file_path: 文件路径
    
    Returns:
        dict: 包含文本内容、文件类型和警告信息
    """
    file_ext = Path(file_path).suffix.lower()
    
    if file_ext == ".pdf":
        return parse_pdf(file_path)
    elif file_ext == ".docx":
        return parse_docx(file_path)
    elif file_ext in [".txt", ".md"]:
        return parse_text(file_path)
    else:
        return {
            "text": "",
            "file_type": file_ext,
            "warnings": ["不支持的文件格式"]
        }

def parse_pdf(file_path: str) -> dict:
    """解析PDF文件
    
    Args:
        file_path: PDF文件路径
    
    Returns:
        dict: 包含文本内容、文件类型和警告信息
    """
    try:
        import pypdfium2 as pdfium
        
        text = ""
        warnings = []
        
        # 打开PDF文件
        pdf = pdfium.PdfDocument(file_path)
        
        # 提取每页文本
        for page_num in range(len(pdf)):
            page = pdf.get_page(page_num)
            text += page.get_text()
        
        pdf.close()
        
        # 检查是否为空文本（可能是扫描版PDF）
        if not text.strip():
            warnings.append("扫描版 PDF 可能无法完整解析")
        
        return {
            "text": text,
            "file_type": "pdf",
            "warnings": warnings
        }
    except Exception as e:
        return {
            "text": "",
            "file_type": "pdf",
            "warnings": [f"PDF解析失败: {str(e)}"]
        }

def parse_docx(file_path: str) -> dict:
    """解析DOCX文件
    
    Args:
        file_path: DOCX文件路径
    
    Returns:
        dict: 包含文本内容、文件类型和警告信息
    """
    try:
        from docx import Document
        
        doc = Document(file_path)
        text = ""
        
        # 提取段落文本
        for para in doc.paragraphs:
            text += para.text + "\n"
        
        # 提取表格文本
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + "\t"
                text += "\n"
        
        return {
            "text": text,
            "file_type": "docx",
            "warnings": []
        }
    except Exception as e:
        return {
            "text": "",
            "file_type": "docx",
            "warnings": [f"DOCX解析失败: {str(e)}"]
        }

def parse_text(file_path: str) -> dict:
    """解析TXT或MD文件
    
    Args:
        file_path: 文本文件路径
    
    Returns:
        dict: 包含文本内容、文件类型和警告信息
    """
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        
        file_ext = Path(file_path).suffix.lower()
        return {
            "text": text,
            "file_type": "txt" if file_ext == ".txt" else "md",
            "warnings": []
        }
    except Exception as e:
        return {
            "text": "",
            "file_type": "txt",
            "warnings": [f"文本文件解析失败: {str(e)}"]
        }